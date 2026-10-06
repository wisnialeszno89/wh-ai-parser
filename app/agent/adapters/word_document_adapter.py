from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Iterable


@dataclass(frozen=True)
class WordDocumentSummary:
    """Model-safe description of a Word document."""

    path: str
    paragraph_count: int
    table_count: int

    def to_payload(self) -> dict[str, object]:
        return {
            "path": self.path,
            "paragraph_count": self.paragraph_count,
            "table_count": self.table_count,
        }


class WordDocumentAdapter:
    """Semantic .docx document adapter with explicit filesystem roots.

    This adapter works with the Word document format directly rather than
    replaying fragile UI actions. It is intentionally conservative: it
    creates and reads documents and appends paragraphs, while path access
    remains restricted to configured roots.

    Microsoft Word itself is not required for these operations. A future
    desktop adapter can layer on top when live Word UI control is needed.
    """

    adapter_id = "word_document"

    def __init__(
        self,
        *,
        allowed_roots: Iterable[str | Path],
    ) -> None:
        roots = tuple(
            Path(root).expanduser().resolve()
            for root in allowed_roots
        )

        if not roots:
            raise ValueError(
                "At least one allowed Word document root is required."
            )

        self._roots = roots

    @property
    def descriptor(self):
        from app.agent.adapters.application_adapter import (
            AdapterDescriptor,
        )

        return AdapterDescriptor(
            adapter_id=self.adapter_id,
            application="Word",
            capabilities=(
                "exists",
                "read_text",
                "create",
                "append_paragraph",
            ),
            description=(
                "Semantic creation and inspection of .docx documents "
                "inside explicitly allowed filesystem roots."
            ),
        )

    def is_available(self) -> bool:
        if not all(
            root.exists() and root.is_dir()
            for root in self._roots
        ):
            return False

        try:
            import importlib.util

            return importlib.util.find_spec("docx") is not None
        except (ImportError, ValueError):
            return False

    @property
    def allowed_roots(self) -> tuple[Path, ...]:
        return self._roots

    def exists(self, path: str | Path) -> bool:
        return self._resolve_for_access(path).exists()

    def read_text(
        self,
        path: str | Path,
    ) -> str:
        document_path = self._resolve_for_access(path)
        self._require_docx(document_path)

        document = self._document(document_path)

        sections: list[str] = []

        for paragraph in document.paragraphs:
            text = paragraph.text.strip()
            if text:
                sections.append(text)

        for table in document.tables:
            for row in table.rows:
                cells = [
                    cell.text.strip()
                    for cell in row.cells
                ]
                if any(cells):
                    sections.append(" | ".join(cells))

        return "\n".join(sections)

    def inspect(
        self,
        path: str | Path,
    ) -> WordDocumentSummary:
        document_path = self._resolve_for_access(path)
        self._require_docx(document_path)

        document = self._document(document_path)

        return WordDocumentSummary(
            path=str(document_path),
            paragraph_count=len(document.paragraphs),
            table_count=len(document.tables),
        )

    def create(
        self,
        path: str | Path,
        *,
        title: str | None = None,
        paragraphs: Iterable[str] = (),
        overwrite: bool = False,
    ) -> Path:
        document_path = self._resolve_for_create(path)
        self._require_docx(document_path)

        if document_path.exists() and not overwrite:
            raise FileExistsError(str(document_path))

        document_path.parent.mkdir(
            parents=True,
            exist_ok=True,
        )

        document = self._document_class()()

        if isinstance(title, str) and title.strip():
            document.add_heading(
                title.strip(),
                level=1,
            )

        for paragraph in paragraphs:
            if not isinstance(paragraph, str):
                raise TypeError(
                    "Word document paragraphs must be strings."
                )
            document.add_paragraph(paragraph)

        temporary = document_path.with_name(
            f".{document_path.name}.agent-tmp.docx"
        )
        document.save(str(temporary))
        temporary.replace(document_path)

        return document_path

    def append_paragraph(
        self,
        path: str | Path,
        text: str,
    ) -> WordDocumentSummary:
        document_path = self._resolve_for_access(path)
        self._require_docx(document_path)

        if not document_path.exists():
            raise FileNotFoundError(str(document_path))

        if not isinstance(text, str):
            raise TypeError(
                "Word paragraph text must be a string."
            )

        document = self._document(document_path)
        document.add_paragraph(text)

        temporary = document_path.with_name(
            f".{document_path.name}.agent-tmp.docx"
        )
        document.save(str(temporary))
        temporary.replace(document_path)

        return WordDocumentSummary(
            path=str(document_path),
            paragraph_count=len(document.paragraphs),
            table_count=len(document.tables),
        )

    def _resolve_for_access(
        self,
        path: str | Path,
    ) -> Path:
        resolved = Path(path).expanduser().resolve(
            strict=False
        )
        self._assert_under_allowed_root(resolved)
        return resolved

    def _resolve_for_create(
        self,
        path: str | Path,
    ) -> Path:
        resolved = Path(path).expanduser().resolve(
            strict=False
        )
        self._assert_under_allowed_root(resolved)

        if resolved.suffix.casefold() != ".docx":
            raise ValueError(
                "WordDocumentAdapter requires a .docx path."
            )

        return resolved

    @staticmethod
    def _require_docx(path: Path) -> None:
        if path.suffix.casefold() != ".docx":
            raise ValueError(
                f"WordDocumentAdapter requires a .docx path, got '{path}'."
            )

    def _assert_under_allowed_root(
        self,
        path: Path,
    ) -> None:
        for root in self._roots:
            try:
                path.relative_to(root)
                return
            except ValueError:
                continue

        raise PermissionError(
            f"Path '{path}' is outside all allowed Word document roots."
        )

    @staticmethod
    def _document_class():
        try:
            from docx import Document
        except ImportError as exc:
            raise RuntimeError(
                "python-docx is required for Word document operations."
            ) from exc

        return Document

    @classmethod
    def _document(cls, path: Path):
        return cls._document_class()(str(path))
