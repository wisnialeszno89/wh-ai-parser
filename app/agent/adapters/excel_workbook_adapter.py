from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Iterable, Sequence


@dataclass(frozen=True)
class ExcelWorkbookSummary:
    """Model-safe description of one Excel workbook."""

    path: str
    sheet_names: tuple[str, ...]
    max_rows: int
    max_columns: int

    def to_payload(self) -> dict[str, object]:
        return {
            "path": self.path,
            "sheet_names": self.sheet_names,
            "max_rows": self.max_rows,
            "max_columns": self.max_columns,
        }


@dataclass(frozen=True)
class ExcelRange:
    """Semantic rectangular spreadsheet data."""

    sheet: str
    cell_range: str
    values: tuple[tuple[object, ...], ...]

    def to_payload(self) -> dict[str, object]:
        return {
            "sheet": self.sheet,
            "cell_range": self.cell_range,
            "values": [
                list(row)
                for row in self.values
            ],
        }


class ExcelWorkbookAdapter(ApplicationAdapter):
    """Semantic .xlsx workbook adapter with explicit filesystem roots.

    The adapter operates on workbook structure and cell values rather than
    spreadsheet UI coordinates. It is intentionally limited to .xlsx files
    in configured roots; live Excel desktop control remains a separate layer.
    """

    adapter_id = "excel_workbook"

    def __init__(
        self,
        *,
        allowed_roots: Iterable[str | Path],
    ) -> None:
        roots = tuple(
            Path(root).expanduser().resolve()
            for root in allowed_roots
        )

        if not roots:
            raise ValueError(
                "At least one allowed Excel workbook root is required."
            )

        self._roots = roots

    @property
    def descriptor(self) -> AdapterDescriptor:
        return AdapterDescriptor(
            adapter_id=self.adapter_id,
            application="Excel",
            capabilities=(
                "exists",
                "inspect",
                "read_range",
                "write_range",
                "append_rows",
                "create_sheet",
            ),
            description=(
                "Semantic inspection and editing of .xlsx workbooks "
                "inside explicitly allowed filesystem roots."
            ),
        )

    def is_available(self) -> bool:
        if not all(
            root.exists() and root.is_dir()
            for root in self._roots
        ):
            return False

        try:
            import importlib.util

            return importlib.util.find_spec("openpyxl") is not None
        except (ImportError, ValueError):
            return False

    @property
    def allowed_roots(self) -> tuple[Path, ...]:
        return self._roots

    def exists(self, path: str | Path) -> bool:
        return self._resolve_for_access(path).exists()

    def inspect(
        self,
        path: str | Path,
    ) -> ExcelWorkbookSummary:
        workbook_path = self._resolve_for_existing_workbook(path)
        workbook = self._load_workbook(workbook_path)

        max_rows = 0
        max_columns = 0

        for sheet in workbook.worksheets:
            max_rows = max(max_rows, sheet.max_row)
            max_columns = max(max_columns, sheet.max_column)

        return ExcelWorkbookSummary(
            path=str(workbook_path),
            sheet_names=tuple(
                workbook.sheetnames
            ),
            max_rows=max_rows,
            max_columns=max_columns,
        )

    def read_range(
        self,
        path: str | Path,
        *,
        sheet: str,
        cell_range: str,
    ) -> ExcelRange:
        workbook_path = self._resolve_for_existing_workbook(path)
        workbook = self._load_workbook(workbook_path)
        worksheet = self._worksheet(workbook, sheet)

        rows = worksheet[cell_range]
        if hasattr(rows, "value"):
            rows = ((rows,),)

        values = tuple(
            tuple(
                self._model_safe_value(cell.value)
                for cell in row
            )
            for row in rows
        )

        return ExcelRange(
            sheet=worksheet.title,
            cell_range=cell_range,
            values=values,
        )

    def write_range(
        self,
        path: str | Path,
        *,
        sheet: str,
        start_cell: str,
        values: Sequence[Sequence[object]],
    ) -> ExcelRange:
        workbook_path = self._resolve_for_existing_workbook(path)
        workbook = self._load_workbook(workbook_path)
        worksheet = self._worksheet(workbook, sheet)

        rows = self._normalize_values(values)
        self._validate_start_cell(start_cell)

        for row_offset, row in enumerate(rows):
            for column_offset, value in enumerate(row):
                worksheet.cell(
                    row=worksheet[start_cell].row + row_offset,
                    column=worksheet[start_cell].column + column_offset,
                    value=value,
                )

        end_row = worksheet[start_cell].row + len(rows) - 1
        end_column = (
            worksheet[start_cell].column
            + max(len(row) for row in rows)
            - 1
        )

        cell_range = self._range_from_bounds(
            worksheet[start_cell].row,
            worksheet[start_cell].column,
            end_row,
            end_column,
        )

        self._atomic_save(
            workbook,
            workbook_path,
        )

        return self.read_range(
            workbook_path,
            sheet=sheet,
            cell_range=cell_range,
        )

    def append_rows(
        self,
        path: str | Path,
        *,
        sheet: str,
        rows: Sequence[Sequence[object]],
    ) -> ExcelRange:
        workbook_path = self._resolve_for_existing_workbook(path)
        workbook = self._load_workbook(workbook_path)
        worksheet = self._worksheet(workbook, sheet)

        normalized_rows = self._normalize_values(rows)

        start_row = worksheet.max_row + 1
        for row in normalized_rows:
            worksheet.append(list(row))

        end_row = start_row + len(normalized_rows) - 1
        end_column = max(
            len(row)
            for row in normalized_rows
        )

        cell_range = self._range_from_bounds(
            start_row,
            1,
            end_row,
            end_column,
        )

        self._atomic_save(
            workbook,
            workbook_path,
        )

        return self.read_range(
            workbook_path,
            sheet=sheet,
            cell_range=cell_range,
        )

    def create_sheet(
        self,
        path: str | Path,
        *,
        sheet: str,
    ) -> ExcelWorkbookSummary:
        workbook_path = self._resolve_for_existing_workbook(path)
        if not isinstance(sheet, str) or not sheet.strip():
            raise ValueError(
                "Excel sheet name must not be empty."
            )

        workbook = self._load_workbook(workbook_path)
        if sheet in workbook.sheetnames:
            raise ValueError(
                f"Excel sheet '{sheet}' already exists."
            )

        workbook.create_sheet(sheet.strip())
        self._atomic_save(
            workbook,
            workbook_path,
        )

        return self.inspect(workbook_path)

    def create(
        self,
        path: str | Path,
        *,
        sheets: Sequence[str] = ("Sheet1",),
        overwrite: bool = False,
    ) -> ExcelWorkbookSummary:
        workbook_path = self._resolve_for_create(path)

        if workbook_path.exists() and not overwrite:
            raise FileExistsError(str(workbook_path))

        normalized_sheets = tuple(
            item.strip()
            for item in sheets
            if isinstance(item, str) and item.strip()
        )

        if not normalized_sheets:
            raise ValueError(
                "At least one Excel sheet name is required."
            )

        if len(set(normalized_sheets)) != len(normalized_sheets):
            raise ValueError(
                "Excel sheet names must be unique."
            )

        workbook = self._workbook_class()()
        default_sheet = workbook.active
        default_sheet.title = normalized_sheets[0]

        for sheet in normalized_sheets[1:]:
            workbook.create_sheet(sheet)

        workbook_path.parent.mkdir(
            parents=True,
            exist_ok=True,
        )

        self._atomic_save(
            workbook,
            workbook_path,
        )

        return self.inspect(workbook_path)

    def _resolve_for_access(
        self,
        path: str | Path,
    ) -> Path:
        resolved = Path(path).expanduser().resolve(
            strict=False
        )
        self._assert_under_allowed_root(resolved)
        return resolved

    def _resolve_for_existing_workbook(
        self,
        path: str | Path,
    ) -> Path:
        resolved = self._resolve_for_access(path)

        if resolved.suffix.casefold() != ".xlsx":
            raise ValueError(
                "ExcelWorkbookAdapter requires a .xlsx path."
            )
        if not resolved.exists():
            raise FileNotFoundError(str(resolved))
        if not resolved.is_file():
            raise IsADirectoryError(str(resolved))

        return resolved

    def _resolve_for_create(
        self,
        path: str | Path,
    ) -> Path:
        resolved = self._resolve_for_access(path)

        if resolved.suffix.casefold() != ".xlsx":
            raise ValueError(
                "ExcelWorkbookAdapter requires a .xlsx path."
            )

        return resolved

    def _worksheet(
        self,
        workbook,
        sheet: str,
    ):
        if not isinstance(sheet, str) or not sheet.strip():
            raise ValueError(
                "Excel sheet name must not be empty."
            )

        if sheet not in workbook.sheetnames:
            raise KeyError(
                f"Excel sheet '{sheet}' does not exist."
            )

        return workbook[sheet]

    @staticmethod
    def _normalize_values(
        values: Sequence[Sequence[object]],
    ) -> tuple[tuple[object, ...], ...]:
        if not isinstance(values, Sequence) or isinstance(
            values,
            (str, bytes),
        ):
            raise TypeError(
                "Excel values must be a sequence of rows."
            )

        rows = tuple(
            tuple(row)
            for row in values
        )

        if not rows:
            raise ValueError(
                "Excel values must not be empty."
            )

        if any(
            len(row) == 0
            for row in rows
        ):
            raise ValueError(
                "Excel rows must not be empty."
            )

        return rows

    @staticmethod
    def _validate_start_cell(cell: str) -> None:
        try:
            from openpyxl.utils.cell import coordinate_from_string

            coordinate_from_string(cell)
        except (ImportError, TypeError, ValueError):
            raise ValueError(
                f"Invalid Excel start cell '{cell}'."
            )

    @staticmethod
    def _range_from_bounds(
        start_row: int,
        start_column: int,
        end_row: int,
        end_column: int,
    ) -> str:
        try:
            from openpyxl.utils import get_column_letter
        except ImportError as exc:
            raise RuntimeError(
                "openpyxl is required for Excel workbook operations."
            ) from exc

        start = (
            f"{get_column_letter(start_column)}{start_row}"
        )
        end = (
            f"{get_column_letter(end_column)}{end_row}"
        )

        return (
            start
            if start == end
            else f"{start}:{end}"
        )

    @staticmethod
    def _model_safe_value(value: object) -> object:
        if value is None or isinstance(
            value,
            (str, int, float, bool),
        ):
            return value

        if hasattr(value, "isoformat"):
            try:
                return value.isoformat()
            except (AttributeError, TypeError, ValueError):
                pass

        return str(value)

    @staticmethod
    def _workbook_class():
        try:
            from openpyxl import Workbook
        except ImportError as exc:
            raise RuntimeError(
                "openpyxl is required for Excel workbook operations."
            ) from exc

        return Workbook

    @staticmethod
    def _load_workbook(path: Path):
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise RuntimeError(
                "openpyxl is required for Excel workbook operations."
            ) from exc

        return load_workbook(
            filename=str(path),
            data_only=False,
        )

    @staticmethod
    def _atomic_save(workbook, path: Path) -> None:
        temporary = path.with_name(
            f".{path.name}.agent-tmp.xlsx"
        )
        workbook.save(str(temporary))
        temporary.replace(path)

    def _assert_under_allowed_root(
        self,
        path: Path,
    ) -> None:
        for root in self._roots:
            try:
                path.relative_to(root)
                return
            except ValueError:
                continue

        raise PermissionError(
            f"Path '{path}' is outside all allowed Excel workbook roots."
        )


from app.agent.adapters.application_adapter import (
    AdapterDescriptor,
    ApplicationAdapter,
)
